import re, sys
sys.path.insert(0,'/mnt/user-data/outputs'); from pp_rows import data_rows
from openpyxl import load_workbook
X='/mnt/user-data/outputs/PantryPulse-MASTER_2026-08-17_EOD.xlsx'; D='2026-09-18'
def set_items(wb, sheet, items, note='18-Sep count (operator)'):
    w=wb[sheet]; log=[]
    for tup in items:
        pat,q,u,cat,disp=tup[0],tup[1],tup[2],tup[3],tup[4]; exc=tup[5] if len(tup)>5 else None
        r=None
        for i in data_rows(w):
            n=w.cell(i,2).value
            if not n or not re.search(pat,str(n),re.I): continue
            if exc and re.search(exc,str(n),re.I): continue
            if isinstance(w.cell(i,3).value,(int,float)): r=i;break
        if r is None:
            r=w.max_row+1
            for c,v in [(1,cat),(2,disp),(3,0),(4,u),(8,1)]: w.cell(r,c).value=v
        old=w.cell(r,3).value; ou=w.cell(r,4).value
        w.cell(r,2).value=disp; w.cell(r,1).value=cat; w.cell(r,3).value=q; w.cell(r,4).value=u; w.cell(r,5).value='in' if q>0 else 'out'; w.cell(r,7).value=None; w.cell(r,9).value=D; w.cell(r,10).value=D; w.cell(r,11).value=note
        log.append(f'{disp}: {old}{ou or ""}->{q}{u}')
    return log
