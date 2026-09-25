"""pp_snapshot — freeze every household's board for a given date into the 'Board Snapshots' sheet.
Called automatically by pp_git_push.py on every push (snapshots the CLOSED day, HOV_CLOSED_DATE),
so the last push of a day = that day's locked board. Idempotent: re-running for the same date replaces it.

Columns: date | household | hid | category | item | qty | unit | status
"""
import re, datetime, sys
from openpyxl import load_workbook
sys.path.insert(0,'/mnt/user-data/outputs')
from pp_rows import data_rows
COLS=['date','household','hid','category','item','qty','unit','status']
def snapshot(xlsx_path, date_iso, replace=True):
    wb=load_workbook(xlsx_path)
    ix=wb['Index']
    houses=[(str(ix.cell(r,1).value),str(ix.cell(r,3).value)) for r in range(5,ix.max_row+1) if ix.cell(r,1).value]
    if 'Board Snapshots' not in wb.sheetnames:
        s=wb.create_sheet('Board Snapshots'); s.append(COLS); s.freeze_panes='A2'
    s=wb['Board Snapshots']
    if replace:
        # drop existing rows for this date (bottom-up)
        for r in range(s.max_row,1,-1):
            if str(s.cell(r,1).value or '')[:10]==date_iso: s.delete_rows(r)
    n=0
    for sheet,hid in houses:
        if sheet not in wb.sheetnames: continue
        w=wb[sheet]
        for r in data_rows(w):
            item=w.cell(r,2).value
            if not item: continue
            q=w.cell(r,3).value
            if not isinstance(q,(int,float)) or q<=0: continue   # snapshot = what is ON the board
            s.append([date_iso,sheet,hid,str(w.cell(r,1).value or ''),str(item),q,str(w.cell(r,4).value or ''),str(w.cell(r,5).value or 'in')]); n+=1
    wb.save(xlsx_path)
    return n
if __name__=='__main__':
    import sys
    p=sys.argv[1]; d=sys.argv[2] if len(sys.argv)>2 else datetime.date.today().isoformat()
    print('snapshot rows written:',snapshot(p,d),'for',d)
