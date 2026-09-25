import re, sys
sys.path.insert(0,"/mnt/user-data/outputs"); from pp_rows import data_rows
from openpyxl import load_workbook
X='/mnt/user-data/outputs/PantryPulse-MASTER_2026-08-17_EOD.xlsx'
D='2026-09-19'
PERISH={'Vegetable','Fruit','Meat','Meat & fish','Dairy','Bread & bakery','Frozen','Herb','Herbs'}
def norm(s): return re.sub(r'[^a-z0-9]','',str(s).lower())
def qparse(s):
    s=str(s).lower().replace('1/4','0.25').replace('1/2','0.5').replace('kg','000g').replace('1l','1000ml').replace('1 ltr','1000ml')
    m=re.search(r'([\d.]+)\s*(ml|g|gm|pc|pcs|loaf|slice|slices|pack)?',s)
    if not m: return None,None
    v=float(m.group(1));u=(m.group(2) or 'pc')
    if u=='gm':u='g'
    if u in('pcs',):u='pc'
    if u=='slices':u='slice'
    return v,u
# items: list of (regex, qty, unit, category, display, [exc])
def apply_overwrite(wb, sheet, items, remove_unlisted=True, protect=None):
    w=wb[sheet]; protect=protect or []
    set_rows=set()
    log=[]
    for tup in items:
        pat,q,u,cat,disp = tup[0],tup[1],tup[2],tup[3],tup[4]
        exc=tup[5] if len(tup)>5 else None
        r=None
        for i in data_rows(w):
            n=w.cell(i,2).value
            if not n or not re.search(pat,str(n),re.I): continue
            if exc and re.search(exc,str(n),re.I): continue
            cur=w.cell(i,3).value
            if isinstance(cur,(int,float)): r=i;break
        if r is None:
            r=w.max_row+1
            for c,v in [(1,cat),(2,disp),(3,0),(4,u),(8,1)]: w.cell(r,c).value=v
        # unit collision: setting grams on a pc row (or vice versa) -> just overwrite to the new unit (it's a SET, not add)
        w.cell(r,2).value=disp; w.cell(r,1).value=cat; w.cell(r,3).value=q; w.cell(r,4).value=u; w.cell(r,5).value='in' if q>0 else 'out'; w.cell(r,7).value=None; w.cell(r,9).value=D; w.cell(r,10).value=D; w.cell(r,11).value='16-Sep perishable overwrite (operator)'
        set_rows.add(r); log.append(f'{disp}={q}{u}')
    removed=[]
    if remove_unlisted:
        for r in data_rows(w):
            if r in set_rows: continue
            n=w.cell(r,2).value; c=str(w.cell(r,1).value or '')
            if n and c in PERISH and isinstance(w.cell(r,3).value,(int,float)) and w.cell(r,3).value>0:
                if any(re.search(p,str(n),re.I) for p in protect): continue
                removed.append(f'{n} ({w.cell(r,3).value}{w.cell(r,4).value or ""})')
                w.cell(r,3).value=0;w.cell(r,5).value='out';w.cell(r,9).value=D;w.cell(r,11).value='16-Sep: removed — not in perishable count'
    return log, removed
