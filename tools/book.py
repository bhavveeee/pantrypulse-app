import sys; sys.path.insert(0,'/mnt/user-data/outputs')
from pp_changelog import log as _cl
from openpyxl import load_workbook
import re
X='/mnt/user-data/outputs/PantryPulse-MASTER_2026-08-17_EOD.xlsx'
def book(sheet,hid,items,D,note):
    wb=load_workbook(X); w=wb[sheet]; orders=wb['Orders']
    used=set(); log=[]
    for it in items:
        pat,q,u=it[0],it[1],it[2]
        mode=it[3] if len(it)>3 else 'add'
        cat =it[4] if len(it)>4 else 'Vegetable'
        nm  =it[5] if len(it)>5 else None
        exc =it[6] if len(it)>6 else None
        r=None; fb=None
        for i in range(5,w.max_row+1):
            n=w.cell(i,2).value
            if not n or i in used or not re.search(pat,str(n),re.I): continue
            if exc and re.search(exc,str(n),re.I): continue
            cur=w.cell(i,3).value
            if isinstance(cur,(int,float)) and cur>0: r=i; break
            if fb is None: fb=i
        if r is None: r=fb
        if r is None:
            r=w.max_row+1
            for c,v in [(1,cat),(2,nm or pat),(3,q),(4,u),(5,'in'),(8,1),(9,D),(10,D),(11,note+' — NEW row')]: w.cell(r,c).value=v
            log.append(f'  +{q}{u:<5} {(nm or pat)[:30]:<30} NEW r{r}')
        else:
            old=w.cell(r,3).value; old=old if isinstance(old,(int,float)) else 0
            nv=q if mode=='set' else round(old+q,2)
            w.cell(r,3).value=nv; w.cell(r,4).value=u; w.cell(r,5).value='in'
            w.cell(r,7).value=None; w.cell(r,9).value=D; w.cell(r,10).value=D; w.cell(r,11).value=note
            log.append(f'  {"=" if mode=="set" else "+"}{q}{u:<5} {str(w.cell(r,2).value)[:30]:<30} {old} -> {nv}')
        orders.append([sheet,hid,D,str(w.cell(r,2).value),q,u,'booked',note])
        _cl(wb,sheet,'order',str(w.cell(r,2).value),(old if 'old' in dir() else None),w.cell(r,3).value,u,'Bhavya-dictated order list',note,'',f'{sheet}!{r}',hid=hid,when=D)
        used.add(r)
    wb.save(X)
    print(f'===== {sheet} ({len(items)} lines) =====')
    [print(x) for x in log]
