import sys; sys.path.insert(0,'/mnt/user-data/outputs')
from pp_changelog import log as _cl
from openpyxl import load_workbook
import re, sys
X='/mnt/user-data/outputs/PantryPulse-MASTER_2026-08-17_EOD.xlsx'
PER={'Vegetable','Fruit','FnV','Herb','Dairy','Meat','Meat & fish','Protein','Bread & bakery','Bakery','Frozen','Produce'}
def canon(s):
    s=str(s or '').lower(); s=re.sub(r'\([^)]*\)',' ',s); s=re.sub(r'[^a-z0-9 ]',' ',s)
    return re.sub(r'\s+',' ',s).strip()
def apply_count(sheet, items, D, note, keep_cats=None):
    """items: (pattern, qty, unit, category_if_new, name_if_new, exclude)"""
    wb=load_workbook(X); w=wb[sheet]
    kept=set(); log={'set':[],'new':[],'rm':[]}
    for it in items:
        pat,q,u = it[0],it[1],it[2]
        cat = it[3] if len(it)>3 else 'Vegetable'
        nm  = it[4] if len(it)>4 else None
        exc = it[5] if len(it)>5 else None
        r=None; fallback=None
        for i in range(5,w.max_row+1):
            n=w.cell(i,2).value
            if not n or not re.search(pat,str(n),re.I): continue
            if exc and re.search(exc,str(n),re.I): continue
            if i in kept: continue
            cur=w.cell(i,3).value
            if isinstance(cur,(int,float)) and cur>0: r=i; break
            if fallback is None: fallback=i
        if r is None: r=fallback
        if r is None:
            r=w.max_row+1
            for c,v in [(1,cat),(2,nm or pat),(3,q),(4,u),(5,'in'),(8,1),(9,D),(10,D),(11,note+' — NEW row')]: w.cell(r,c).value=v
            log['new'].append(f"{(nm or pat)[:32]:<32} {q}{u}  (r{r})")
        else:
            old=w.cell(r,3).value; ou=w.cell(r,4).value
            w.cell(r,3).value=q; w.cell(r,4).value=u; w.cell(r,5).value='in' if q>0 else 'out'
            w.cell(r,7).value=None; w.cell(r,9).value=D; w.cell(r,10).value=D; w.cell(r,11).value=note
            log['set'].append(f"{str(w.cell(r,2).value)[:32]:<32} {old}{ou or ''} -> {q}{u}")
            _cl(wb,sheet,'board-overwrite',str(w.cell(r,2).value),old,q,u,'Bhavya overwrote — smart-list check result',note,'',f'{sheet}!{r}',when=D)
        kept.add(r)
    for i in range(5,w.max_row+1):
        if i in kept: continue
        c=str(w.cell(i,1).value)
        if c not in PER: continue
        if keep_cats and c in keep_cats: continue
        old=w.cell(i,3).value
        if not isinstance(old,(int,float)) or old==0: continue
        w.cell(i,3).value=0; w.cell(i,5).value='out'; w.cell(i,9).value=D; w.cell(i,10).value=D
        w.cell(i,11).value=f"{note}: not in the count (was {old}{w.cell(i,4).value or ''})"
        log['rm'].append(f"{str(w.cell(i,2).value)[:32]:<32} {old}{w.cell(i,4).value or ''}  [{c}]")
        _cl(wb,sheet,'board-overwrite',str(w.cell(i,2).value),old,0,str(w.cell(i,4).value or ''),'Bhavya overwrote — smart-list check result',note+': not in the count','',f'{sheet}!{i}',when=D)
    wb.save(X)
    print(f"===== {sheet} =====")
    print(f"  SET ({len(log['set'])}):");  [print('    '+x) for x in log['set']]
    print(f"  NEW ({len(log['new'])}):");  [print('    '+x) for x in log['new']]
    print(f"  REMOVED ({len(log['rm'])}):"); [print('    '+x) for x in log['rm']]
    return kept
